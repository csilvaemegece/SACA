import base64
import os
import re
import uuid
from datetime import date, datetime
from io import BytesIO

import openpyxl
from flask import (Blueprint, current_app, flash, jsonify, redirect,
                    render_template, request, send_file, send_from_directory,
                    url_for)
from openpyxl.utils import get_column_letter

from app import db
from app.models import MotivoVisita, Registro
from app.ocr_parser import parse_cedula, validar_run

bp = Blueprint("main", __name__)

DATA_URL_RE = re.compile(r"^data:image/(\w+);base64,(.+)$")


@bp.route("/")
def index():
    motivos = MotivoVisita.query.filter_by(activo=True).order_by(MotivoVisita.nombre).all()
    return render_template("index.html", active_page="registrar", motivos=motivos)


@bp.route("/ocr", methods=["POST"])
def ocr():
    data = request.get_json(silent=True) or {}
    image_data = data.get("image", "")
    match = DATA_URL_RE.match(image_data)
    if not match:
        return jsonify({"error": "Formato de imagen inválido"}), 400

    raw = base64.b64decode(match.group(2))
    filename = f"{uuid.uuid4().hex}.jpg"
    path = os.path.join(current_app.config["UPLOAD_FOLDER"], filename)
    with open(path, "wb") as f:
        f.write(raw)

    try:
        fields = parse_cedula(path)
    except Exception:
        current_app.logger.exception("Error al procesar OCR de %s", filename)
        return jsonify({
            "error": "No se pudo leer el documento automáticamente, completa los datos a mano.",
            "foto_filename": filename,
            "foto_url": url_for("main.uploaded_file", filename=filename),
        }), 200

    fields.pop("texto_ocr_crudo", None)
    fields["foto_filename"] = filename
    fields["foto_url"] = url_for("main.uploaded_file", filename=filename)
    return jsonify(fields)


@bp.route("/uploads/<path:filename>")
def uploaded_file(filename):
    return send_from_directory(current_app.config["UPLOAD_FOLDER"], filename)


@bp.route("/registrar", methods=["POST"])
def registrar():
    form = request.form
    nombres = form.get("nombres", "").strip()
    apellidos = form.get("apellidos", "").strip()
    rut = form.get("rut", "").strip()
    fecha_nacimiento_raw = form.get("fecha_nacimiento", "").strip()
    email = form.get("email", "").strip()
    telefono = form.get("telefono", "").strip()
    foto_filename = form.get("foto_filename", "").strip()
    motivo_visita_id_raw = form.get("motivo_visita_id", "").strip()

    errors = []
    if not nombres:
        errors.append("El nombre es obligatorio.")
    if not apellidos:
        errors.append("El apellido es obligatorio.")
    if not rut:
        errors.append("El RUT es obligatorio.")
    elif not validar_run(rut):
        errors.append("El RUT ingresado no es válido, revísalo.")

    motivo = None
    if not motivo_visita_id_raw:
        errors.append("El motivo de la visita es obligatorio.")
    else:
        motivo = MotivoVisita.query.get(motivo_visita_id_raw)
        if motivo is None:
            errors.append("El motivo de la visita seleccionado no es válido.")

    fecha_nacimiento = None
    if not fecha_nacimiento_raw:
        errors.append("La fecha de nacimiento es obligatoria.")
    else:
        try:
            fecha_nacimiento = datetime.strptime(fecha_nacimiento_raw, "%Y-%m-%d").date()
            if fecha_nacimiento > date.today():
                errors.append("La fecha de nacimiento no puede ser futura.")
        except ValueError:
            errors.append("Formato de fecha de nacimiento inválido.")

    if errors:
        for e in errors:
            flash(e, "error")
        foto_url = url_for("main.uploaded_file", filename=foto_filename) if foto_filename else ""
        motivos = MotivoVisita.query.filter_by(activo=True).order_by(MotivoVisita.nombre).all()
        return render_template(
            "index.html",
            form_data=form,
            foto_url=foto_url,
            active_page="registrar",
            motivos=motivos,
        ), 400

    registro = Registro(
        nombres=nombres,
        apellidos=apellidos,
        rut=rut,
        fecha_nacimiento=fecha_nacimiento,
        email=email or None,
        telefono=telefono or None,
        foto_filename=foto_filename or None,
        motivo_visita_id=motivo.id,
    )
    db.session.add(registro)
    db.session.commit()
    flash(f"Registro guardado: {nombres} {apellidos}", "success")
    return redirect(url_for("main.index"))


def _filtered_registros(args):
    query = Registro.query
    desde = args.get("desde", "").strip()
    hasta = args.get("hasta", "").strip()

    if desde:
        try:
            d = datetime.strptime(desde, "%Y-%m-%d").date()
            query = query.filter(Registro.fecha_registro >= datetime.combine(d, datetime.min.time()))
        except ValueError:
            desde = ""
    if hasta:
        try:
            h = datetime.strptime(hasta, "%Y-%m-%d").date()
            query = query.filter(Registro.fecha_registro <= datetime.combine(h, datetime.max.time()))
        except ValueError:
            hasta = ""

    return query.order_by(Registro.fecha_registro.desc()).all(), desde, hasta


@bp.route("/reportes")
def reportes():
    registros, desde, hasta = _filtered_registros(request.args)
    return render_template(
        "reportes.html",
        registros=registros,
        desde=desde,
        hasta=hasta,
        active_page="reportes",
    )


@bp.route("/exportar")
def exportar():
    registros, desde, hasta = _filtered_registros(request.args)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Registros"
    headers = ["Nombres", "Apellidos", "RUT", "Fecha de nacimiento", "Edad",
               "Motivo de la visita", "Email", "Teléfono", "Fecha de registro"]
    ws.append(headers)
    for r in registros:
        ws.append([
            r.nombres,
            r.apellidos,
            r.rut,
            r.fecha_nacimiento.strftime("%d-%m-%Y") if r.fecha_nacimiento else "",
            r.edad,
            r.motivo.nombre if r.motivo else "",
            r.email or "",
            r.telefono or "",
            r.fecha_registro.strftime("%d-%m-%Y %H:%M") if r.fecha_registro else "",
        ])
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(len(header) + 2, 18)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"registros_adultos_mayores_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.route("/configuracion")
def configuracion():
    motivos = MotivoVisita.query.order_by(MotivoVisita.nombre).all()
    return render_template("configuracion.html", motivos=motivos, active_page="configuracion")


@bp.route("/configuracion/motivos", methods=["POST"])
def motivos_agregar():
    nombre = request.form.get("nombre", "").strip()
    if not nombre:
        flash("El nombre del motivo no puede estar vacío.", "error")
    elif MotivoVisita.query.filter(db.func.lower(MotivoVisita.nombre) == nombre.lower()).first():
        flash(f"Ya existe un motivo llamado «{nombre}».", "error")
    else:
        db.session.add(MotivoVisita(nombre=nombre, activo=True))
        db.session.commit()
        flash(f"Motivo «{nombre}» agregado.", "success")
    return redirect(url_for("main.configuracion"))


@bp.route("/configuracion/motivos/<int:motivo_id>/alternar", methods=["POST"])
def motivos_alternar(motivo_id):
    motivo = MotivoVisita.query.get_or_404(motivo_id)
    motivo.activo = not motivo.activo
    db.session.commit()
    estado = "activado" if motivo.activo else "desactivado"
    flash(f"Motivo «{motivo.nombre}» {estado}.", "success")
    return redirect(url_for("main.configuracion"))
